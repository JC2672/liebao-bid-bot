"""Part 2 (queue side): import shortlisted rows, list the queue, and the
per-row actions (Open Folder, Applied, Remove, Retry, Generate).

The SQLite `opportunities` table is a work queue only - see docs/ARCHITECTURE.md.
Generation itself (ChatGPT web -> resume_json -> PDF) lives in generation.py;
this router only ever flips a row to `generating` and enqueues it there - see
generation.py's module docstring for why (a single background worker, not
this request, does the actual work).
"""
from __future__ import annotations

import io
import os
import shutil
import sqlite3
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, HTTPException, UploadFile
from openpyxl import load_workbook

from .. import generation, sheets
from ..db import get_conn
from ..models import BulkIds, ImportDirectRequest, Opportunity
from ..profiles import get_profile

router = APIRouter(prefix="/queue", tags=["queue"])

DATA_DIR = Path(__file__).resolve().parent.parent.parent.parent / "data"
STAGING_DIR = DATA_DIR / "staging"


def _row_to_opportunity(row) -> Opportunity:
    return Opportunity(
        id=row["id"], profile=row["profile"], company=row["company"], title=row["title"],
        location=row["location"], source=row["source"], url=row["url"],
        description=row["description"], status=row["status"], error=row["error"],
        resume_json=row["resume_json"], staging_dir=row["staging_dir"],
        created_at=row["created_at"], updated_at=row["updated_at"],
    )


@router.get("", response_model=list[Opportunity])
def list_queue(profile: str) -> list[Opportunity]:
    with get_conn() as conn:
        rows = conn.execute(
            "SELECT * FROM opportunities WHERE profile = ? ORDER BY created_at ASC", (profile,)
        ).fetchall()
    return [_row_to_opportunity(r) for r in rows]


@router.get("/progress")
def queue_progress(profile: str | None = None) -> dict:
    """Profile-scoped remaining/ETA for the generation run currently in
    flight - see generation.py's get_progress() for the actual tracking.
    `profile` omitted: the floating status widget (App.tsx) isn't scoped
    to any one profile, so it asks "is anything running, and whose" -
    get_progress() resolves which profile that is on its own."""
    return generation.get_progress(profile)


def _applied_urls(prof) -> set[str]:
    try:
        return sheets.get_applied_urls(prof.sheet_webapp_url, prof.sheet_secret, prof.sheet_tab)
    except RuntimeError:
        return set()  # Sheet Web App not configured yet; don't block import


def _dedupe_and_insert(
    conn: sqlite3.Connection,
    profile: str,
    applied_urls: set[str],
    existing_urls: set[str],
    rows: list[tuple[str, str, str, str, str, str]],
) -> tuple[list[int], int]:
    """Shared by both import paths below - the dedupe rule (skip anything
    already applied-to or already queued for this profile, by URL) and the
    insert itself are the one thing that actually matters here; where the
    rows came from (a hand-pruned XLSX upload, or a Fetch run's shortlist
    passed straight through) is the only difference between the two callers."""
    inserted_ids: list[int] = []
    skipped = 0
    for company, title, location, source, url, description in rows:
        url = url.strip()
        if not url or url in applied_urls or url in existing_urls:
            skipped += 1
            continue
        cur = conn.execute(
            """INSERT INTO opportunities (profile, company, title, location, source, url, description)
               VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (profile, company, title, location, source, url, description),
        )
        inserted_ids.append(cur.lastrowid)
        existing_urls.add(url)
    return inserted_ids, skipped


@router.post("/import")
async def import_shortlist(profile: str, file: UploadFile) -> dict:
    """Accepts the XLSX exported by /fetch/export (pruned by hand in Excel),
    dedupes against the profile's Sheet tab + current queue, and inserts the rest."""
    prof = get_profile(profile)  # raises 404-able error if unknown

    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {h: i for i, h in enumerate(headers)}

    applied_urls = _applied_urls(prof)

    with get_conn() as conn:
        existing_urls = {
            r["url"] for r in conn.execute(
                "SELECT url FROM opportunities WHERE profile = ?", (profile,)
            ).fetchall()
        }
        rows = [
            (
                str(row[col["company"]] or ""),
                str(row[col["title"]] or ""),
                str(row[col["location"]] or ""),
                str(row[col["source"]] or ""),
                str(row[col["url"]] or ""),
                str(row[col["description"]] or ""),
            )
            for row in ws.iter_rows(min_row=2, values_only=True)
        ]
        inserted_ids, skipped = _dedupe_and_insert(conn, profile, applied_urls, existing_urls, rows)

    return {"inserted": len(inserted_ids), "skipped": skipped}


@router.post("/import-direct")
def import_direct(body: ImportDirectRequest) -> dict:
    """Same dedupe as /import above, but given a Fetch run's shortlist rows
    directly instead of requiring an export-to-XLSX-then-upload roundtrip -
    powers Fetch's "Queue after fetch" checkbox. Returns the newly-inserted
    rows' ids (not just a count) so the caller can immediately bulk-generate
    exactly those rows, without touching ones that were already queued."""
    prof = get_profile(body.profile)
    applied_urls = _applied_urls(prof)

    with get_conn() as conn:
        existing_urls = {
            r["url"] for r in conn.execute(
                "SELECT url FROM opportunities WHERE profile = ?", (body.profile,)
            ).fetchall()
        }
        rows = [(r.company, r.title, r.location, r.source, r.url, r.description) for r in body.rows]
        inserted_ids, skipped = _dedupe_and_insert(conn, body.profile, applied_urls, existing_urls, rows)

    return {"inserted": len(inserted_ids), "skipped": skipped, "ids": inserted_ids}


@router.post("/{opp_id}/applied")
def mark_applied(opp_id: int) -> dict:
    """Sheet append happens FIRST, before anything on disk moves - on
    purpose. This used to move the files out of staging before appending to
    the Sheet, which meant a Sheet failure (missing/misconfigured Web App,
    network hiccup) left the row stuck: files already relocated to the real
    output folder, but the queue row still `ready` with a `staging_dir` that
    no longer existed, so retrying it could never actually finish. Doing the
    Sheet append first means a failure here leaves everything exactly as it
    was - staging intact, row still `ready` - so it's safe to just fix the
    Sheet config and click Applied again."""
    with get_conn() as conn:
        row = conn.execute("SELECT * FROM opportunities WHERE id = ?", (opp_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "Opportunity not found")
        opp = _row_to_opportunity(row)

    if opp.status != "ready":
        raise HTTPException(400, f"Opportunity is '{opp.status}', not ready to apply")

    prof = get_profile(opp.profile)

    try:
        sheets.append_applied_row(
            prof.sheet_webapp_url, prof.sheet_secret, prof.sheet_tab,
            date=date.today().isoformat(), company=opp.company, title=opp.title,
            source=opp.source, url=opp.url,
        )
    except RuntimeError as exc:
        raise HTTPException(500, f"Sheet append failed, nothing was moved: {exc}") from exc

    safe = lambda s: "".join(c for c in s if c not in '\\/:*?"<>|').strip()
    dest = Path(prof.output_root) / f"{safe(opp.company)} - {safe(opp.title)} - {date.today():%Y-%m-%d}"
    dest.mkdir(parents=True, exist_ok=True)

    if opp.staging_dir and Path(opp.staging_dir).exists():
        for f in Path(opp.staging_dir).iterdir():
            shutil.move(str(f), dest / f.name)
        shutil.rmtree(opp.staging_dir, ignore_errors=True)

    with get_conn() as conn:
        conn.execute("DELETE FROM opportunities WHERE id = ?", (opp_id,))

    return {"folder": str(dest)}


@router.post("/{opp_id}/remove")
def remove_opportunity(opp_id: int) -> dict:
    with get_conn() as conn:
        row = conn.execute("SELECT status, staging_dir FROM opportunities WHERE id = ?", (opp_id,)).fetchone()
        if row is None:
            raise HTTPException(404, "Opportunity not found")
        if row["status"] == "generating":
            # Backend-side safety net, independent of the Queue screen's
            # own lock on this button - generation.py's worker now bails
            # out if a row's status moves away from `generating` before it
            # gets dequeued (see run_generation's early-exit check), but
            # that only covers "hasn't started yet". One already mid-flight
            # through an actual ChatGPT call can't be interrupted; deleting
            # it here wouldn't crash anything (the eventual write just
            # affects zero rows), but it would waste that call for nothing.
            raise HTTPException(409, "Opportunity is actively generating - wait for it to finish first")
        if row["staging_dir"] and Path(row["staging_dir"]).exists():
            shutil.rmtree(row["staging_dir"], ignore_errors=True)
        conn.execute("DELETE FROM opportunities WHERE id = ?", (opp_id,))
    return {"ok": True}


@router.post("/{opp_id}/retry")
def retry_opportunity(opp_id: int) -> dict:
    # Also allowed from `generating` - the escape hatch for a row stuck
    # there with nothing actually working on it anymore (the ChatGPT
    # window was closed mid-run, the backend restarted mid-run, etc.).
    # generation.py's own worker is now hardened against silently dying,
    # but this is the user-facing recovery path regardless of cause.
    with get_conn() as conn:
        cur = conn.execute(
            "UPDATE opportunities SET status = 'queued', error = NULL, updated_at = datetime('now') "
            "WHERE id = ? AND status IN ('failed', 'generating')",
            (opp_id,),
        )
        if cur.rowcount == 0:
            raise HTTPException(400, "Opportunity is not in 'failed' or 'generating' state")
    return {"ok": True}


@router.post("/{opp_id}/generate")
def generate_opportunity(opp_id: int) -> dict:
    """Flips the row to `generating` and hands it to generation.py's
    background worker - returns immediately, the actual ChatGPT/PDF work
    happens off-request. The Queue screen's existing poll picks up the
    eventual ready/failed transition.

    `open_when_ready=True` here (and only here, not bulk-generate below):
    the single per-row Generate button pops the staging folder open in
    Explorer the moment this row lands on `ready`, so there's no need to
    also click Open Folder - a bulk batch doing the same per row would mean
    several Explorer windows popping open back to back instead."""
    with get_conn() as conn:
        cur = conn.execute(
            "UPDATE opportunities SET status = 'generating', error = NULL, updated_at = datetime('now') "
            "WHERE id = ? AND status = 'queued'",
            (opp_id,),
        )
        if cur.rowcount == 0:
            raise HTTPException(400, "Opportunity is not in 'queued' state")
    generation.enqueue(opp_id, open_when_ready=True)
    return {"ok": True}


@router.post("/bulk-remove")
def bulk_remove(body: BulkIds) -> dict:
    if not body.ids:
        return {"removed": 0}
    with get_conn() as conn:
        placeholders = ",".join("?" * len(body.ids))
        # Same "generating" exclusion as the single-row /remove above,
        # just silent here (skip rather than 400) - a bulk action already
        # mixes statuses by nature, so one actively-generating row in the
        # selection shouldn't block removing the rest of it.
        rows = conn.execute(
            f"SELECT id, staging_dir FROM opportunities WHERE id IN ({placeholders}) AND status != 'generating'",
            body.ids,
        ).fetchall()
        ids = [r["id"] for r in rows]
        for row in rows:
            if row["staging_dir"] and Path(row["staging_dir"]).exists():
                shutil.rmtree(row["staging_dir"], ignore_errors=True)
        if ids:
            conn.execute(
                f"DELETE FROM opportunities WHERE id IN ({','.join('?' * len(ids))})", ids
            )
    return {"removed": len(rows)}


@router.post("/bulk-retry")
def bulk_retry(body: BulkIds) -> dict:
    if not body.ids:
        return {"retried": 0}
    with get_conn() as conn:
        cur = conn.execute(
            f"UPDATE opportunities SET status = 'queued', error = NULL, updated_at = datetime('now') "
            f"WHERE id IN ({','.join('?' * len(body.ids))}) AND status IN ('failed', 'generating')",
            body.ids,
        )
    return {"retried": cur.rowcount}


@router.post("/bulk-generate")
def bulk_generate(body: BulkIds) -> dict:
    if not body.ids:
        return {"generating": 0}
    with get_conn() as conn:
        placeholders = ",".join("?" * len(body.ids))
        queued_ids = [
            r["id"] for r in conn.execute(
                f"SELECT id FROM opportunities WHERE id IN ({placeholders}) AND status = 'queued'",
                body.ids,
            ).fetchall()
        ]
        if queued_ids:
            conn.execute(
                f"UPDATE opportunities SET status = 'generating', error = NULL, updated_at = datetime('now') "
                f"WHERE id IN ({','.join('?' * len(queued_ids))})",
                queued_ids,
            )
    for opp_id in queued_ids:
        generation.enqueue(opp_id)
    return {"generating": len(queued_ids)}


@router.post("/{opp_id}/open-folder")
def open_folder(opp_id: int) -> dict:
    """Browsers can't open local folders directly - the backend does it via Explorer."""
    with get_conn() as conn:
        row = conn.execute("SELECT staging_dir FROM opportunities WHERE id = ?", (opp_id,)).fetchone()
    if row is None or not row["staging_dir"]:
        raise HTTPException(404, "No staging folder for this opportunity yet")
    os.startfile(row["staging_dir"])  # Windows-only, matches "My machine only" scope
    return {"ok": True}
